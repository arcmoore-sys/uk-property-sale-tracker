#!/usr/bin/env python3
"""Populate a deal's mandate-tracker.xlsx from its tracker.json.

Usage: python refresh_tracker.py <deal-folder>

Design (v0.6.0): the workbook is the firm's master template copied verbatim into
the deal folder at setup (assets/mandate-tracker-template.xlsx). This script opens
that copied workbook and writes the live state from tracker.json into the existing
tabs IN PLACE, preserving all of the template's formatting, formulas, dropdowns and
layout. It never rebuilds the workbook from code. If the workbook is missing (for
example a skill ran before setup copied it), the template is copied in first.

Skills update tracker.json, then call this. tracker.json stays canonical; the
workbook is the rendered, human-facing view.
"""
import json, os, sys, subprocess, glob
from copy import copy
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))           # .../shared
def plugin_root():
    env = os.environ.get("CLAUDE_PLUGIN_ROOT")
    if env and os.path.isdir(os.path.join(env, "assets")):
        return env
    return os.path.dirname(HERE)
ROOT = plugin_root()
TEMPLATE = os.path.join(ROOT, "assets", "mandate-tracker-template.xlsx")

NDA_MAP = {"not_sent":"Not sent","sent":"Sent","signed":"Signed"}
AML_MAP = {"pending":"Pending","clear":"Clear","flagged":"Flagged"}
POF_MAP = {"pending":"Pending","verified":"Verified","insufficient":"Insufficient"}
BID_MAP = {"none":"None","requested":"Requested","received":"Received","declined":"Declined"}
YESNO   = {True:"Yes",False:"No","yes":"Yes","no":"No","y":"Yes","n":"No",
           "true":"Yes","false":"No","n/a":"N/A","na":"N/A"}

def yn(v):
    if isinstance(v, bool): return YESNO[v]
    if v is None: return ""
    s = str(v).strip()
    return YESNO.get(s.lower(), s)

def norm(s):
    return " ".join(str(s or "").split()).strip().lower()

# ---- styling helpers (preserve the template, extend it consistently) --------
def styled_extent(ws, col=1, start=4, cap=2000):
    """Last contiguous row (from start) whose first column carries a border."""
    last = start - 1
    r = start
    while r <= cap:
        if ws.cell(row=r, column=col).border.left.style is not None:
            last = r; r += 1
        else:
            break
    return last

def copy_row_style(ws, src_row, dst_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(row=src_row, column=c); d = ws.cell(row=dst_row, column=c)
        d.font = copy(s.font); d.fill = copy(s.fill); d.border = copy(s.border)
        d.alignment = copy(s.alignment); d.number_format = s.number_format
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height

def ensure_row_style(ws, r, extent, ncols):
    """For rows past the template's styled band, mirror a parity-matched band row
    (row 4 even / row 5 odd) so zebra striping and borders continue."""
    if r > extent:
        copy_row_style(ws, 4 if r % 2 == 0 else 5, r, ncols)

def setv(ws, r, c, v):
    ws.cell(row=r, column=c).value = v

def clear_values(ws, rows, cols):
    for r in rows:
        for c in cols:
            ws.cell(row=r, column=c).value = None

# ---- per-tab population ------------------------------------------------------
def pop_mandate(ws, deal):
    g = deal.get
    pairs = {
        "B4":g("name",""),  "D4":g("mandate_ref",""),
        "B5":g("address",""),"D5":g("sector",""),
        "B6":g("client",""), "D6":g("tenure",""),
        "B7":g("sale_structure",""),"D7":g("vat_position",""),
        "B9":g("mandate_type",""),  "D9":g("instruction_date",""),
        "B10":g("agent",""),        "D10":g("fee_basis",""),
        "B11":g("agent_email",""),  "D11":g("abort_fee",""),
        "B12":g("process_type",""), "D12":g("send_mode",""),
        "B13":g("quoting_price",""),"D13":g("target_niy",""),
        "B15":g("launch_date",""),  "D15":g("bid_request_date_round_1",""),
        "B16":g("bid_request_date_round_2",""),"D16":g("preferred_bidder_date",""),
        "B17":g("hot_date",""),     "D17":g("target_exchange",""),
        "B18":g("target_completion",""),"D18":g("actual_completion",""),
        "D25":g("current_stage",""),
    }
    for coord, val in pairs.items():
        ws[coord] = val   # value only; formulas in B21:B25/D21:D24 untouched

def pop_keyed(ws, items, key_field, key_col, write_map, ncols, start=4):
    """Update rows that match an existing key cell; append unmatched rows."""
    extent = styled_extent(ws)
    index = {}
    r = start
    while r <= extent:
        kv = ws.cell(row=r, column=key_col).value
        if kv not in (None, ""):
            index[norm(kv)] = r
        r += 1
    next_row = extent + 1
    for it in items:
        k = norm(it.get(key_field, ""))
        if not k:
            continue
        if k in index:
            row = index[k]
        else:
            row = next_row; next_row += 1
            ensure_row_style(ws, row, extent, ncols)
            setv(ws, row, key_col, it.get(key_field, ""))
        for col, fn in write_map.items():
            val = fn(it)
            if val is not None:
                setv(ws, row, col, val)

def pop_buyers(ws, buyers):
    extent = styled_extent(ws); ncols = 15
    clear_values(ws, range(4, extent + 1), range(1, ncols + 1))
    for i, b in enumerate(buyers):
        r = 4 + i
        ensure_row_style(ws, r, extent, ncols)
        setv(ws,r,1,b.get("name","")); setv(ws,r,2,b.get("company","")); setv(ws,r,3,b.get("email",""))
        setv(ws,r,4,b.get("enquiry_date","")); setv(ws,r,5,NDA_MAP.get(b.get("nda_status",""),""))
        setv(ws,r,6,b.get("nda_sent_date","")); setv(ws,r,7,b.get("last_chase_date",""))
        setv(ws,r,8,b.get("chase_count","")); setv(ws,r,9,b.get("nda_signed_date",""))
        setv(ws,r,10,AML_MAP.get(b.get("aml_status",""),"")); setv(ws,r,11,POF_MAP.get(b.get("pof_status",""),""))
        setv(ws,r,12,BID_MAP.get(b.get("bid_status_round_1",""),"")); setv(ws,r,13,BID_MAP.get(b.get("bid_status_round_2",""),""))
        setv(ws,r,14,BID_MAP.get(b.get("bid_status_round_3",""),"")); setv(ws,r,15,b.get("notes",""))

def pop_bidlog(ws, deal, bids):
    from datetime import date as _date
    cr = int(deal.get("current_round", 0) or 0) or 1
    ws["A2"] = f"{deal.get('name','(property)')}:  Round {cr} bid comparison        Prepared {_date.today().strftime('%d %b %Y')}"
    GREEN, RED, GREY = "1F7A3D", "C0392B", "777777"
    by = {}
    for bd in bids:
        try: rd = int(bd.get("round",1) or 1)
        except (TypeError, ValueError): rd = 1
        by.setdefault(bd.get("buyer",""), {})[rd] = bd
    rows = []
    for buyer, rounds in by.items():
        if not buyer: continue
        present = sorted(rounds); latest_r = present[-1]; latest = rounds[latest_r]
        lp = latest.get("offer_price")
        if len(present) == 1:
            mv, mc = ("Round 1", GREY) if present[0]==1 else (f"New in R{present[0]}", GREEN)
        else:
            a = rounds[present[-2]].get("offer_price"); b = rounds[present[-1]].get("offer_price")
            if isinstance(a,(int,float)) and isinstance(b,(int,float)):
                mv, mc = ("Improved", GREEN) if b>a else (("Reduced", RED) if b<a else ("Held", GREY))
            else: mv, mc = "Updated", GREY
        if cr >= 2 and latest_r < cr:
            mv, mc = f"Withdrawn (last R{latest_r})", RED
        rows.append({"buyer":buyer,"rounds":rounds,"latest":latest,
                     "lp": lp if isinstance(lp,(int,float)) else -1,"mv":mv,"mc":mc})
    rows.sort(key=lambda x: x["lp"], reverse=True)

    extent = styled_extent(ws)            # data band (notes sit below it)
    band = list(range(4, extent + 1))
    clear_values(ws, band, [1,2,3,4,8,9,10,11,12,13,14,15])   # keep E,F,G formulas
    from openpyxl.styles import Font
    for i, rec in enumerate(rows):
        r = 4 + i
        if r > extent: break
        leader = (i == 0 and rec["lp"] >= 0)
        setv(ws,r,1,rec["buyer"])
        for col, rd in ((2,1),(3,2),(4,3)):
            v = rec["rounds"].get(rd, {}).get("offer_price")
            setv(ws,r,col, v if isinstance(v,(int,float)) else None)
        for col, rd in ((8,1),(9,2),(10,3)):
            v = rec["rounds"].get(rd, {}).get("niy")
            setv(ws,r,col, v if isinstance(v,(int,float)) else None)
        km = ws.cell(row=r, column=11); km.value = rec["mv"]
        km.font = Font(name=km.font.name, size=km.font.size, bold=True, color=rec["mc"])
        lt = rec["latest"]
        setv(ws,r,12,lt.get("conditions","")); setv(ws,r,13,lt.get("completion",""))
        setv(ws,r,14,lt.get("pof","") or lt.get("funding_type","")); setv(ws,r,15,lt.get("notes",""))
        if leader:
            cg = ws.cell(row=r, column=7)
            cg.font = Font(name=cg.font.name, size=cg.font.size, bold=True, color="B98A2E")

def populate(folder, data):
    out = os.path.join(folder, "mandate-tracker.xlsx")
    if not os.path.exists(out):
        import shutil
        shutil.copy2(TEMPLATE, out)        # pull the firm's template, do not rebuild
        os.chmod(out, 0o644)               # copy2 preserves the template's read-only bits; make the deal copy writable
    wb = openpyxl.load_workbook(out)
    deal = data.get("deal", {})

    if "Mandate Summary" in wb.sheetnames:
        pop_mandate(wb["Mandate Summary"], deal)
    if "Task Checklist" in wb.sheetnames:
        pop_keyed(wb["Task Checklist"], data.get("tasks", []), "task", 2, {
            3: lambda it: it.get("owner") or None,
            4: lambda it: it.get("status") or None,
            5: lambda it: it.get("due") or None,
            6: lambda it: it.get("done") or None,
            7: lambda it: it.get("automated_by") or None,
        }, ncols=7)
    if "Information & Data Room" in wb.sheetnames:
        pop_keyed(wb["Information & Data Room"], data.get("dataroom", []), "document", 2, {
            1: lambda it: it.get("category") or None,
            3: lambda it: yn(it.get("received")) or None,
            4: lambda it: yn(it.get("in_dataroom")) or None,
            5: lambda it: it.get("date") or None,
            6: lambda it: it.get("notes") or None,
        }, ncols=6)
    if "Buyer Pipeline" in wb.sheetnames:
        pop_buyers(wb["Buyer Pipeline"], data.get("buyers", []))
    if "Bid Log" in wb.sheetnames:
        pop_bidlog(wb["Bid Log"], deal, data.get("bids", []))
    if "Legal & Completion" in wb.sheetnames:
        pop_keyed(wb["Legal & Completion"], data.get("legals", []), "milestone", 1, {
            2: lambda it: it.get("responsible") or None,
            3: lambda it: it.get("status") or None,
            4: lambda it: it.get("target_date") or None,
            5: lambda it: it.get("actual_date") or None,
            6: lambda it: it.get("notes") or None,
        }, ncols=6)
    if "Client Reporting" in wb.sheetnames:
        ws = wb["Client Reporting"]; extent = styled_extent(ws); ncols = 6
        clear_values(ws, range(4, extent + 1), range(1, ncols + 1))
        for i, it in enumerate(data.get("reporting", [])):
            r = 4 + i; ensure_row_style(ws, r, extent, ncols)
            setv(ws,r,1,it.get("date","")); setv(ws,r,2,it.get("type","")); setv(ws,r,3,it.get("period",""))
            setv(ws,r,4,it.get("sent_to","")); setv(ws,r,5,it.get("method","")); setv(ws,r,6,it.get("notes",""))
    if os.path.exists(out):
        os.chmod(out, 0o644)               # workbook may have been copied read-only (copy2 preserves template bits); ensure writable before saving in place
    wb.save(out)
    return out

def find_recalc():
    for p in ["/mnt/skills/public/xlsx/scripts/recalc.py",
              os.path.expanduser("~/.claude/skills/xlsx/scripts/recalc.py")]:
        if os.path.exists(p):
            return p
    hits = glob.glob("/mnt/skills/**/xlsx/scripts/recalc.py", recursive=True)
    return hits[0] if hits else None

def main():
    if len(sys.argv) < 2:
        print("usage: refresh_tracker.py <deal-folder>"); sys.exit(1)
    folder = sys.argv[1]
    with open(os.path.join(folder, "tracker.json")) as f:
        data = json.load(f)
    out = populate(folder, data)
    rc = find_recalc()
    if rc:
        try:
            subprocess.run([sys.executable, rc, out], capture_output=True, timeout=120)
        except Exception:
            pass
    print(f"refreshed {out}")

if __name__ == "__main__":
    main()
